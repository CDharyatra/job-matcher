from pathlib import Path

import pytest

from jobagent.resume import build_profile

FIX = Path(__file__).parent / "fixtures"


@pytest.mark.skipif(not (FIX / "sample_resume.docx").exists(), reason="no docx fixture")
def test_docx_parse():
    p = build_profile([str(FIX / "sample_resume.docx")])
    assert "Senior Backend Engineer" in p.headline
    assert p.total_years == 6.0
    for skill in ("python", "postgresql", "aws", "kubernetes"):
        assert skill in p.skills
    assert p.bullets


@pytest.mark.skipif(not (FIX / "sample_resume.pdf").exists(), reason="no pdf fixture")
def test_pdf_parse():
    p = build_profile([str(FIX / "sample_resume.pdf")])
    assert p.total_years == 6.0
    for skill in ("python", "postgresql", "aws"):
        assert skill in p.skills


def test_genai_skills_extracted(tmp_path):
    import docx
    doc = docx.Document()
    doc.add_paragraph("Machine Learning Engineer with 3 years of experience.")
    doc.add_paragraph("Built RAG pipelines with LangChain, LangGraph, pgvector and AWS Bedrock.")
    doc.add_paragraph("Semantic search over OpenSearch with vector embeddings; Ollama for local eval.")
    path = tmp_path / "genai.docx"
    doc.save(str(path))
    p = build_profile([str(path)])
    for skill in ("rag", "langchain", "langgraph", "pgvector", "bedrock",
                  "semantic search", "vector embeddings", "opensearch", "ollama",
                  "machine learning"):
        assert skill in p.skills, skill


def test_cache_key_includes_parser_version(tmp_path, monkeypatch):
    import docx
    import jobagent.resume as rmod

    monkeypatch.setattr(rmod, "CACHE_DIR", tmp_path)
    doc = docx.Document()
    doc.add_paragraph("Python developer with 3 years of experience.")
    path = tmp_path / "r.docx"
    doc.save(str(path))
    first = rmod.load_or_build([str(path)])
    # Simulate an old cache entry written by a previous parser version and a
    # vocab change: bumping the version must invalidate the stale entry.
    monkeypatch.setattr(rmod, "PROFILE_CACHE_VERSION", "v-test-bump")
    second = rmod.load_or_build([str(path)])
    assert first.skills == second.skills
    cache_files = list(tmp_path.glob("resume_profile_*.json"))
    assert len(cache_files) == 2  # old-version entry left untouched, new one written


def test_empty_path_does_not_crash():
    p = build_profile([""])
    assert p.raw_text == ""
    assert p.skills == []